"""
预处理：PPTX → Markdown

这是整个 RAG 的第一道关，也是决定检索效果上限的一步。
PPT 抽出来的原始文本是碎片化的（标题、要点散落在不同 shape 里），
直接向量化会导致语义不完整、检索命中率低。

本脚本做三件事：
  1. 按「章节标题 + 正文」的结构抽文本
  2. 保留来源信息（文件名、章节），写入 Markdown front-matter 风格注释
  3. 输出到 data_processed/，供后续分块使用

用法：
    python preprocess.py              # 处理全部 PPT
    python preprocess.py --inspect    # 只看抽取效果，不写文件
"""

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation

from config import DEFAULT_CONFIG


def extract_from_shape(shape) -> list:
    """
    从单个 shape 抽取文本行。

    PPT 的文本可能藏在多种 shape 里：
      - 普通文本框 / 占位符 → shape.text_frame
      - 表格                → shape.table
      - 组合形状            → shape.shapes（递归）
    """
    lines = []

    # 组合形状：递归下钻
    if shape.shape_type == 6 or hasattr(shape, "shapes"):
        try:
            for sub in shape.shapes:
                lines.extend(extract_from_shape(sub))
            return lines
        except Exception:
            pass

    # 表格
    if getattr(shape, "has_table", False):
        try:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append("| " + " | ".join(cells) + " |")
            return lines
        except Exception:
            pass

    # 普通文本框
    if getattr(shape, "has_text_frame", False):
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip()
            if not text:
                continue
            # 按缩进层级推断列表/标题
            level = para.level or 0
            if level > 0:
                lines.append("  " * (level - 1) + "- " + text)
            else:
                lines.append(text)

    return lines


def extract_from_pptx(path: Path) -> list:
    """
    抽取一整个 PPTX，返回 [(slide_no, [line, ...]), ...]
    """
    prs = Presentation(str(path))
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            lines.extend(extract_from_shape(shape))
        # 去重连续重复行（PPT 里常见标题被重复放置）
        cleaned = []
        for ln in lines:
            if not cleaned or cleaned[-1] != ln:
                cleaned.append(ln)
        if cleaned:
            slides.append((idx, cleaned))
    return slides


def slides_to_markdown(path: Path, slides: list) -> str:
    """
    把抽取结果组织成 Markdown。

    结构：
        # <文件名>
        <!-- source: xxx.pptx -->
        ## 第 N 页
        ...内容...
    """
    stem = path.stem
    out = [f"# {stem}", f"<!-- source: {path.name} -->", ""]

    for slide_no, lines in slides:
        out.append(f"## 第 {slide_no} 页")
        out.append("")
        # 首个短行大概率是标题，提升为 ###
        for i, ln in enumerate(lines):
            if i == 0 and len(ln) <= 30 and not ln.startswith(("-", "|", " ")):
                out.append(f"### {ln}")
            else:
                out.append(ln)
        out.append("")

    text = "\n".join(out)
    # 压掉 3 个以上连续空行
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def find_pptx(data_dir: Path) -> list:
    """按章节号排序找出所有 PPTX"""
    files = list(data_dir.glob("*.pptx")) + list(data_dir.glob("*.ppt"))

    def sort_key(p: Path):
        m = re.search(r"第\s*(\d+)\s*章", p.name)
        return (int(m.group(1)) if m else 999, p.name)

    return sorted(files, key=sort_key)


def main():
    parser = argparse.ArgumentParser(description="PPTX → Markdown 预处理")
    parser.add_argument("--inspect", action="store_true",
                        help="只打印抽取效果，不写文件")
    args = parser.parse_args()

    cfg = DEFAULT_CONFIG
    data_dir = Path(cfg.data_dir)

    if not data_dir.exists():
        print(f"[错误] 素材目录不存在：{data_dir}")
        return 1

    files = find_pptx(data_dir)
    if not files:
        print(f"[错误] 在 {data_dir} 中没找到 .pptx 文件")
        return 1

    print(f"找到 {len(files)} 个 PPTX 文件\n")

    if not args.inspect:
        cfg.ensure_dirs()
        out_dir = Path(cfg.processed_dir)

    total_chars = 0
    for f in files:
        try:
            slides = extract_from_pptx(f)
        except Exception as e:
            print(f"[跳过] {f.name} —— 解析失败：{e}")
            continue

        md = slides_to_markdown(f, slides)
        total_chars += len(md)

        print(f"[{'预览' if args.inspect else '输出'}] {f.name}")
        print(f"        幻灯片 {len(slides)} 页 → {len(md)} 字符")

        if args.inspect:
            # 打印前 25 行，用来判断抽取质量
            preview = md.splitlines()[:25]
            print("        " + "-" * 50)
            for ln in preview:
                print("        | " + ln)
            print("        " + "-" * 50)
        else:
            out_path = out_dir / (f.stem + ".md")
            out_path.write_text(md, encoding="utf-8")

        print()

    print(f"合计约 {total_chars} 字符")
    if not args.inspect:
        print(f"已输出到：{cfg.processed_dir}")
        print("\n下一步：python -m unittest 或 python build_index.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
